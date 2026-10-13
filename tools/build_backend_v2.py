from datetime import date
from pathlib import Path

from docx import Document

from build_backend_doc import (
    add_body,
    add_bullet,
    add_callout_paragraph,
    add_code,
    add_heading,
    add_number,
    add_table,
    configure_document,
)


ROOT = Path(r"D:\yxy5.0")
OUTPUT = ROOT / "docs" / "照片记忆本后端技术方案_V2.docx"


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = 38
    title.paragraph_format.space_after = 15
    title.add_run("照片记忆本后端技术方案")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = 22
    subtitle.add_run("V2  长期副本与总结能力技术设计")
    add_body(doc, "本文档基于《照片记忆本产品需求文档 PRD V2》，面向后端开发、前端接入和后续部署，给出用户、记忆、主题、长期副本、图片处理、月度年度总结和阶段总结的完整技术方案。")
    add_callout_paragraph(doc, "推荐结论  ", "采用 NestJS 模块化单体作为 API 服务，使用 PostgreSQL 保存业务数据，使用 S3 兼容对象存储保存图片，使用 Redis 和 BullMQ 承载 EXIF、缩略图、周期总结和 AI 生成任务。")
    add_table(doc, ["项目", "技术决策"], [
        ["架构", "模块化单体，按领域拆分 NestJS Module"],
        ["API", "RESTful API + OpenAPI，统一返回格式"],
        ["数据库", "PostgreSQL + Prisma"],
        ["图片存储", "S3 / MinIO，原图和缩略图分离"],
        ["异步任务", "Redis + BullMQ，支持重试和幂等"],
        ["总结生成", "规则统计 + 可替换 AI Summary Provider"],
        ["部署", "Docker + Nginx，后续接入 HTTPS、备份和监控"],
    ], widths=[1.35, 5.1])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = 12
    p.add_run(f"版本  V2.0    编制日期  {date.today().isoformat()}")

    doc.add_page_break()
    add_heading(doc, "一 建设目标与范围", 1)
    add_body(doc, "V2 后端需要支撑四类核心数据：图片记忆、当日状态、当日主题和长期副本。用户先上传图片，再补充备忘录、状态和主题；当主题与已有长期副本匹配时，系统自动建立关联。总结功能在用户授权后读取指定数据，先生成预览，用户确认后保存为不可变快照。")
    add_heading(doc, "P0 范围", 2)
    for item in [
        "用户注册、登录、Token 刷新和数据隔离。",
        "单图记忆创建、编辑、详情和删除；空字段不产生待补充入口。",
        "当日状态多选和当日主题多选。",
        "主题创建、查询和记忆关联。",
        "长期副本创建、生命周期管理和时间轴查询。",
        "图片 EXIF 时间解析、缩略图和签名访问地址。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "P1 范围", 2)
    for item in [
        "月度总结和年度总结。",
        "长期副本 30 天、100 天和一周年阶段总结。",
        "总结授权、照片/文字数据开关和预览确认。",
        "规则统计、代表照片挑选、AI 文案和总结快照。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "非目标", 2)
    add_body(doc, "本阶段不实现多人协作、公开分享、地图轨迹、人物关系、多媒体混剪和复杂通知。数据模型为未来多图和云同步预留扩展，但不让这些能力阻塞核心交付。")

    add_heading(doc, "二 总体架构", 1)
    add_code(doc, "React Web\n    │ HTTPS / REST / OpenAPI\n    ▼\nNestJS API ───────── PostgreSQL\n    │                      │\n    ├── S3 / MinIO         └── Prisma migrations\n    │\n    ├── Redis + BullMQ ─── Worker\n    │                       ├── EXIF / 缩略图\n    │                       ├── 周期总结统计\n    │                       └── AI 文案生成\n    │\n    └── Summary Provider（规则模板 / AI Provider 可替换）")
    add_body(doc, "采用模块化单体而不是微服务，原因是当前业务规模和团队规模更适合集中交付。模块之间通过 Service 接口协作，Controller 不直接访问 Prisma；异步任务通过队列隔离图片处理和总结生成，避免用户接口长时间等待。")
    add_heading(doc, "技术栈", 2)
    add_table(doc, ["层次", "技术", "用途"], [
        ["API 服务", "NestJS + TypeScript", "认证、业务模块、参数校验和接口文档"],
        ["数据库", "PostgreSQL", "用户、记忆、主题、副本和总结快照"],
        ["数据访问", "Prisma", "类型安全 Schema、迁移和事务"],
        ["对象存储", "S3 / MinIO", "原图、缩略图、总结封面"],
        ["任务系统", "Redis + BullMQ", "图片处理、周期任务、AI 生成和重试"],
        ["观测", "Pino + Prometheus", "结构化日志、指标和错误定位"],
    ], widths=[1.25, 2.2, 3.0])

    add_heading(doc, "三 领域模型与数据库设计", 1)
    add_heading(doc, "Memory 记忆", 2)
    add_code(doc, "Memory {\n  id, userId, note, eventAt, timezone\n  statuses[], isPending, createdAt, updatedAt, deletedAt\n}\nMemoryMedia { memoryId, mediaId }\nMemoryTopic { memoryId, topicId }")
    add_body(doc, "一条记忆必须关联一张图片，可选备忘录、多个当日状态和多个当日主题。eventAt 是时间线排序依据；createdAt 只表示写入时间。isPending 由服务端根据 note、statuses 和 topics 重新计算。")
    add_heading(doc, "Topic 主题", 2)
    add_code(doc, "Topic { id, userId, name, color, icon, isNone, createdAt, updatedAt }\nunique(userId, name)")
    add_body(doc, "主题是记忆和长期副本之间的连接点。系统预置“无”，用户可以创建恋爱、考研、读研、旅行等自定义主题。“无”不参与长期副本关联。")
    add_heading(doc, "LongTermCopy 长期副本", 2)
    add_code(doc, "LongTermCopy {\n  id, userId, name, topicId, startDate\n  coverMediaId, lifecycleStatus\n  createdAt, updatedAt, completedAt, archivedAt\n}\nLongTermCopyMemory { copyId, memoryId, linkedAt }")
    add_body(doc, "一个长期副本只关联一个主题。一条记忆可以选择多个主题，因此可以进入多个长期副本。LongTermCopyMemory 作为显式关联表保存，方便固定时间轴结果、执行删除清理和生成总结。")
    add_heading(doc, "MediaAsset 图片资源", 2)
    add_code(doc, "MediaAsset {\n  id, userId, objectKey, thumbnailKey, originalName\n  mimeType, fileSize, width, height\n  captureAt, captureSource, processingStatus, createdAt\n}")
    add_heading(doc, "Summary 数据模型", 2)
    add_code(doc, "SummaryConsent { userId, enabled, allowPhoto, allowText, version, grantedAt, revokedAt }\nSummaryJob { id, userId, scopeType, scopeId, periodStart, periodEnd, status, idempotencyKey }\nSummaryDraft { id, jobId, cardsJson, selectedMediaIds, generatedText, createdAt }\nSummarySnapshot { id, userId, scopeType, scopeId, periodStart, periodEnd, cardsJson, version, confirmedAt }")
    add_body(doc, "Draft 和 Snapshot 分离：Draft 允许用户替换照片、删除卡片和修改文案；确认后写入 Snapshot，后续记忆变化不修改历史快照。")
    add_heading(doc, "重要日子", 2)
    add_code(doc, "ImportantDay { id, userId, name, date, createdAt, updatedAt }\nImportantDayMemory { importantDayId, memoryId, createdAt }\nunique(importantDayId, memoryId)")

    doc.add_page_break()
    add_heading(doc, "四 图片上传与记忆创建", 1)
    add_code(doc, "选择图片\n  → POST /api/media/upload-url\n  → 前端直接 PUT 到 S3 / MinIO\n  → POST /api/media/:id/complete\n  → Worker 解析 EXIF / 生成缩略图\n  → POST /api/memories 创建记忆\n  → 按 topics[] 关联 LongTermCopy")
    add_heading(doc, "图片规则", 2)
    for item in [
        "支持 JPEG、PNG、WebP 和 HEIC；服务端检查文件头和真实 MIME。",
        "单张图片限制建议为 20MB，用户和空间维度增加配额。",
        "objectKey 使用随机 ID，不使用用户原文件名作为存储路径。",
        "原图和缩略图分别保存，图片下载使用短时效签名 URL。",
        "Worker 失败时保留原图，记录失败状态并允许重试。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "时间解析优先级", 2)
    add_table(doc, ["优先级", "来源", "captureSource"], [
        ["1", "EXIF DateTimeOriginal", "exif"],
        ["2", "EXIF CreateDate", "exif"],
        ["3", "用户手动填写的事件时间", "manual"],
        ["4", "服务器上传时间", "upload_time"],
    ], widths=[0.75, 3.8, 1.95])
    add_heading(doc, "记忆创建事务", 2)
    for item in [
        "校验当前用户拥有 mediaId 和所有 topicId。",
        "校验图片处理状态为 completed，读取服务端解析的 EXIF。",
        "按 eventAt、timezone、note、statuses 和 topics 写入 Memory。",
        "计算 isPending，并写入 MemoryTopic 关联。",
        "根据 topics 查询用户的进行中 LongTermCopy，创建幂等的 LongTermCopyMemory。",
        "返回包含图片签名 URL、状态和主题的记忆详情。",
    ]:
        add_number(doc, item)
    add_callout_paragraph(doc, "幂等要求  ", "关联操作使用 (copyId, memoryId) 联合唯一键和请求幂等键，重复保存或网络重试不能创建重复节点。")

    add_heading(doc, "五 长期副本与时间轴", 1)
    add_heading(doc, "创建流程", 2)
    for item in [
        "用户进入长期副本页，若没有数据，返回空状态“暂无长期副本”。",
        "用户填写副本名称、关联主题、开始日期和可选封面。",
        "服务端校验主题属于当前用户，写入 LongTermCopy。",
        "系统询问是否导入开始日期之后已有的同主题记忆。",
        "用户确认导入时，按主题和时间范围批量写入 LongTermCopyMemory。",
    ]:
        add_number(doc, item)
    add_heading(doc, "生命周期", 2)
    add_table(doc, ["状态", "服务端规则"], [
        ["进行中", "允许关联新记忆，参与阶段总结任务"],
        ["暂停", "保留数据，默认不触发新阶段总结，可手动恢复"],
        ["已完成", "保留历史，停止自动阶段总结"],
        ["已归档", "从主列表隐藏，仍可查看和恢复"],
    ], widths=[1.35, 5.1])
    add_heading(doc, "时间轴查询", 2)
    add_code(doc, "GET /api/long-term-copies/:id/timeline\n  ?from=2026-01-01&to=2026-12-31\n  &cursor=xxx&pageSize=30")
    add_body(doc, "接口返回记忆节点、事件时间、图片缩略图、备忘录摘要、当日状态和主题。前端可据此实现横向滑动、分页加载和阶段节点定位。查询必须按 userId 和 copyId 同时校验归属。")

    add_heading(doc, "六 API 接口设计", 1)
    add_heading(doc, "认证", 2)
    add_table(doc, ["方法", "路径", "说明"], [
        ["POST", "/api/auth/register", "注册"],
        ["POST", "/api/auth/login", "登录并返回 Token"],
        ["POST", "/api/auth/refresh", "刷新 Access Token"],
        ["POST", "/api/auth/logout", "撤销 Refresh Token"],
        ["GET", "/api/me", "读取当前用户"],
    ], widths=[0.75, 3.0, 2.6])
    add_heading(doc, "记忆、主题和副本", 2)
    add_table(doc, ["方法", "路径", "说明"], [
        ["POST", "/api/media/upload-url", "申请上传地址"],
        ["POST", "/api/media/:id/complete", "确认上传完成"],
        ["POST", "/api/memories", "创建记忆"],
        ["GET", "/api/memories", "列表、搜索和筛选"],
        ["PATCH", "/api/memories/:id", "编辑备忘录、状态、主题和时间"],
        ["DELETE", "/api/memories/:id", "软删除记忆"],
        ["GET", "/api/topics", "读取主题"],
        ["POST", "/api/topics", "新增自定义主题"],
        ["POST", "/api/long-term-copies", "创建长期副本"],
        ["GET", "/api/long-term-copies", "读取副本列表"],
        ["PATCH", "/api/long-term-copies/:id", "更新生命周期或封面"],
        ["GET", "/api/long-term-copies/:id/timeline", "读取时间轴节点"],
    ], widths=[0.75, 3.3, 2.3])
    add_heading(doc, "记忆请求示例", 2)
    add_code(doc, "POST /api/memories\n{\n  \"mediaId\": \"media_xxx\",\n  \"note\": \"今天第一次完成模拟考试。\",\n  \"eventAt\": \"2026-10-12T18:30:00+08:00\",\n  \"statuses\": [\"焦虑\", \"学习\"],\n  \"topics\": [\"topic_exam\"]\n}")

    add_heading(doc, "七 总结系统设计", 1)
    add_heading(doc, "总结范围", 2)
    add_table(doc, ["scopeType", "范围", "生成节点"], [
        ["memory_month", "全部记忆的自然月", "每月 1 日生成上月"],
        ["memory_year", "全部记忆的自然年", "每年 1 月 1 日生成上年"],
        ["copy_milestone", "单个长期副本", "第 30 天、第 100 天、一周年"],
    ], widths=[1.55, 3.0, 2.0])
    add_heading(doc, "授权模型", 2)
    add_code(doc, "SummaryConsent {\n  enabled: boolean,\n  allowPhoto: boolean,\n  allowText: boolean,\n  allowStatus: boolean,\n  allowTopic: boolean,\n  version, grantedAt, revokedAt\n}")
    add_body(doc, "首次生成总结前，如果没有有效授权，API 返回 CONSENT_REQUIRED，前端展示授权页。默认允许照片、文字、心情、状态和主题；用户可以关闭照片或文字。撤回授权后停止自动任务，但不删除已经保存的 SummarySnapshot。")
    add_heading(doc, "生成流水线", 2)
    add_code(doc, "周期调度器\n  → 检查 consent\n  → 创建 SummaryJob（幂等）\n  → 查询授权范围内的 Memory\n  → 规则统计和代表照片挑选\n  → Summary Provider 生成文案\n  → 写入 SummaryDraft\n  → 用户预览、编辑、确认\n  → 写入不可变 SummarySnapshot")
    add_heading(doc, "规则统计", 2)
    for item in [
        "记忆数量、活跃天数和记录最密集日期。",
        "状态频次、主题频次和时间段分布。",
        "代表照片候选按时间覆盖、主题覆盖和图片质量排序。",
        "长期副本已过去天数、节点数量和阶段变化。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "AI 文案 Provider", 2)
    add_body(doc, "AI 只接收经过授权和脱敏后的摘要数据，默认不直接上传所有原始照片。需要原图理解时，使用短时效签名 URL，并在任务完成后失效。Provider 采用接口抽象，支持规则模板、云端模型或本地模型切换。")
    add_code(doc, "interface SummaryProvider {\n  generate(input: SummaryInput): Promise<SummaryText>;\n}\n\nSummaryInput = { stats, selectedQuotes, topics, moodTrend, allowText, allowPhoto }")
    add_heading(doc, "草稿与快照", 2)
    for item in [
        "Draft 支持修改卡片顺序、替换照片、删除卡片和修改文案。",
        "Confirm 接口使用版本号和幂等键，确保重复点击不会保存两份快照。",
        "Snapshot 保存 periodStart、periodEnd、scopeId、输入数据版本和生成版本。",
        "后续数据变化不更新 Snapshot；重新生成会生成新的 version。",
    ]:
        add_number(doc, item)

    doc.add_page_break()
    add_heading(doc, "八 异步任务与调度", 1)
    add_table(doc, ["队列", "任务", "重试策略"], [
        ["media-processing", "EXIF、尺寸、缩略图", "失败重试 3 次，保留原图"],
        ["copy-backfill", "创建副本时导入历史同主题记忆", "批量任务可断点续跑"],
        ["summary-periodic", "月度、年度和阶段总结草稿", "按幂等键避免重复"],
        ["summary-ai", "AI 文案生成", "超时重试，失败回退规则模板"],
        ["media-cleanup", "软删除文件清理", "延迟执行，支持恢复窗口"],
    ], widths=[1.55, 3.35, 1.65])
    add_heading(doc, "幂等键规则", 2)
    add_code(doc, "memory:{userId}:{memoryId}:{updatedAt}\nsummary:{scopeType}:{scopeId}:{periodStart}:{periodEnd}:{consentVersion}\nlink:{copyId}:{memoryId}")
    add_body(doc, "队列任务必须记录 jobId、attempts、lastError 和完成时间。对周期总结而言，同一 scope、周期和授权版本只允许存在一个进行中的任务。")

    add_heading(doc, "九 搜索、索引与性能", 1)
    add_table(doc, ["查询场景", "实现建议"], [
        ["按事件时间排序", "Memory(userId, eventAt DESC) 复合索引"],
        ["按主题筛选", "MemoryTopic(topicId, memoryId) 联合索引"],
        ["长期副本时间轴", "LongTermCopyMemory(copyId, linkedAt) 联合索引"],
        ["标题和备忘录搜索", "初期 ILIKE，规模增长后 PostgreSQL Full Text Search"],
        ["状态和主题统计", "按总结任务范围批量聚合，避免逐条调用 API"],
    ], widths=[2.2, 4.25])
    add_body(doc, "列表接口采用 cursor 分页，避免时间线数据增加后 offset 越来越慢。图片列表只返回缩略图签名 URL，详情页再按需返回原图签名 URL。")

    add_heading(doc, "十 安全、隐私与数据治理", 1)
    add_table(doc, ["风险", "控制措施"], [
        ["跨用户读取", "所有查询带 userId，详情、主题、副本和媒体都做归属校验"],
        ["图片泄露", "对象存储私有桶，使用短时效签名 URL"],
        ["总结越权", "AI 任务只读取授权范围内的字段和媒体"],
        ["Token 泄露", "Access Token 短时效，Refresh Token 可撤销并轮换"],
        ["恶意上传", "文件头检测、大小限制、配额、限流和异步解码"],
        ["误删除", "软删除和恢复窗口，清理任务延迟执行"],
        ["撤回授权", "立即停止自动生成，不修改已保存快照"],
    ], widths=[1.5, 4.95])
    add_heading(doc, "十一 部署与目录结构", 1)
    add_code(doc, "server/\n├─ src/\n│  ├─ auth/\n│  ├─ users/\n│  ├─ media/\n│  ├─ memories/\n│  ├─ topics/\n│  ├─ long-term-copies/\n│  ├─ summaries/\n│  ├─ important-days/\n│  ├─ common/\n│  └─ main.ts\n├─ prisma/schema.prisma\n├─ workers/\n├─ test/\n└─ docker-compose.yml")
    add_body(doc, "开发环境使用 PostgreSQL、Redis 和 MinIO；测试环境增加真实对象存储兼容性测试；生产环境使用 Nginx、HTTPS、数据库备份、对象存储生命周期和任务监控。")

    add_heading(doc, "十二 前端接入与迁移", 1)
    add_body(doc, "当前前端已经通过 Repository 层访问 IndexedDB。后续只需增加 ApiMemoryRepository、ApiTopicRepository、ApiLongTermCopyRepository 和 ApiSummaryRepository，页面组件不直接依赖 IndexedDB 或 fetch。")
    add_code(doc, "interface MemoryRepository {\n  list(filter?: MemoryFilter): Promise<Memory[]>;\n  get(id: string): Promise<Memory>;\n  create(input: CreateMemoryInput): Promise<Memory>;\n  update(id: string, input: UpdateMemoryInput): Promise<Memory>;\n  remove(id: string): Promise<void>;\n}")
    add_heading(doc, "十三 分阶段实施计划", 1)
    add_table(doc, ["阶段", "主要交付", "验收重点"], [
        ["后端 P0", "认证、媒体、记忆、主题、基础搜索", "单条记忆字段完整且数据隔离"],
        ["副本 P0", "长期副本、主题关联、时间轴、生命周期", "一条记忆可进入多个副本"],
        ["总结 P1", "授权、统计、代表照片、Draft 和 Snapshot", "月度/年度/阶段总结可预览确认"],
        ["AI P1", "Summary Provider、AI 文案、失败回退", "授权范围内生成且可追踪版本"],
        ["平台 P2", "云同步、导入导出、多图、分享", "不改变已有 API 核心模型"],
    ], widths=[1.35, 3.0, 2.1])
    add_heading(doc, "十四 后端验收清单", 1)
    for item in [
        "用户只能读取和修改自己的图片、记忆、主题和长期副本。",
        "一条记忆可以保存图片、备忘录、多个状态和多个主题。",
        "选择主题后，记忆可以幂等地进入对应长期副本。",
        "创建副本时可以按开始日期导入历史同主题记忆。",
        "副本支持进行中、暂停、已完成和已归档。",
        "时间轴能够按 cursor 分页返回图片、文字、状态和主题。",
        "总结生成前必须有有效授权，照片和文字开关生效。",
        "月度、年度和长期副本阶段总结均能生成 Draft。",
        "用户确认后保存 Snapshot，后续记忆变化不修改历史版本。",
        "AI 失败时可回退到规则模板，任务状态和错误原因可追踪。",
    ]:
        add_bullet(doc, item)
    add_callout_paragraph(doc, "实施原则  ", "先保证记忆、主题和副本关系准确，再增加总结和 AI。任何自动化功能都必须可解释、可撤回、可重试，并且不越过用户授权边界。")

    doc.core_properties.title = "照片记忆本后端技术方案 V2"
    doc.core_properties.subject = "长期副本、主题关联和总结能力后端设计"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "照片记忆本, 后端, 长期副本, 总结, NestJS, PostgreSQL"
    doc.save(OUTPUT)
    print(str(OUTPUT))


if __name__ == "__main__":
    build()
