from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from build_backend_doc import (
    add_body, add_bullet, add_callout_paragraph, add_code, add_heading,
    add_number, add_run, add_table, configure_document, set_run_font,
    style_paragraph,
)

OUTPUT = Path(r"D:\yxy5.0\docs\照片记忆本前后端技术方案.docx")


def set_footer(doc, label):
    paragraph = doc.sections[0].footer.paragraphs[0]
    paragraph.clear(); paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(paragraph, f"{label}  |  ", size=8, color="6E7B77")
    run = paragraph.add_run(); set_run_font(run, size=8, color="6E7B77")
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText"); text.set(qn("xml:space"), "preserve"); text.text = "PAGE"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(); configure_document(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    set_footer(doc, "照片记忆本前后端技术方案")

    p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(p, before=48, after=14, line=1.05)
    add_run(p, "照片记忆本前后端技术方案", size=27, bold=True, color="000000")
    p = doc.add_paragraph(); style_paragraph(p, after=22, line=1.25)
    add_run(p, "网页版 MVP 本地 PostgreSQL 实施版", size=13, bold=True, color="315A5C")
    add_body(doc, "本文档对应《照片记忆本产品需求文档 V2》，描述当前代码可落地的前端、后端、数据库、接口、业务事务、测试和部署方案。所有项目代码与文档统一保存在 D:\\yxy5.0。", after=14)
    add_callout_paragraph(doc, "架构结论  ", "MVP 采用 React + TypeScript + Vite 前端、NestJS + Prisma 后端和 PostgreSQL 数据库；图片以 multipart 上传到后端本地目录，数据库保存媒体元数据。当前为单用户本地版本，同时保留 userId 隔离字段。")
    add_table(doc, ["层级", "技术选型", "职责"], [
        ["Web 前端", "React、TypeScript、Vite", "表单、列表、批量任务、授权交互、响应式页面"],
        ["API 后端", "NestJS、TypeScript", "校验、事务、业务规则、媒体访问"],
        ["数据访问", "Prisma ORM", "模型、迁移、事务和 PostgreSQL 查询"],
        ["数据库", "PostgreSQL 18", "记忆、主题、长期副本和关联数据"],
        ["图片存储", "本地 uploads 目录", "MVP 保存原图，生产环境可替换对象存储"],
        ["测试", "Vitest", "前端领域逻辑、组件渲染、后端服务单元测试"],
    ], widths=[1.15, 1.95, 3.35])

    doc.add_page_break()
    add_heading(doc, "1 总体架构", 1)
    add_body(doc, "浏览器通过 Vite 开发服务器访问 React 应用，/api 请求由开发代理转发至 NestJS。后端使用 Prisma 连接 PostgreSQL；图片上传后写入 D:\\yxy5.0\\data\\uploads，Memory 与 MediaAsset 保存业务信息和文件元数据。前端 Repository 层屏蔽 HTTP 细节，使视图组件不直接拼装接口请求。")
    add_code(doc, "Browser  ->  React/Vite :5173  ->  /api proxy  ->  NestJS :3001\n                                                   |\n                                                   +-> Prisma -> PostgreSQL :5433\n                                                   +-> D:\\yxy5.0\\data\\uploads")
    add_heading(doc, "1.1 MVP 边界", 2)
    for item in [
        "本地单用户模式使用固定 LOCAL_USER_ID，但所有核心表保留 userId。",
        "图片存储使用本地目录，不依赖云服务；生产化时切换 S3 兼容存储。",
        "月度、年度和阶段总结仅保留架构扩展位，本次不实现生成任务。",
        "当前 API 为 REST，不引入 GraphQL、消息队列和微服务。",
    ]: add_bullet(doc, item)

    add_heading(doc, "2 前端方案", 1)
    add_heading(doc, "2.1 模块划分", 2)
    add_table(doc, ["模块", "主要文件", "职责"], [
        ["应用编排", "src/App.tsx", "视图切换、数据加载、授权决策和全局交互"],
        ["记忆编辑", "MemoryForm.tsx", "单图表单、多图检测、自定义标签与编辑"],
        ["批量导入", "BatchImportDialog.tsx", "统一主题、队列、进度、失败项和重试"],
        ["副本决策", "LongTermCopyDecisionDialog.tsx", "按主题收集启用、跳过、抑制决策"],
        ["长期副本", "LongTermCopiesView 等", "创建、自动收集开关和时间线"],
        ["数据访问", "src/repositories", "Memory、Topic、LongTermCopy API 封装"],
        ["领域逻辑", "src/domain", "搜索、候选主题判断、批量任务构造"],
    ], widths=[1.2, 2.75, 2.5])
    add_heading(doc, "2.2 状态与一致性", 2)
    add_body(doc, "记忆主列表使用 useReducer 管理加载、单条更新、批量更新、删除和选中状态；主题、长期副本、弹窗和批量任务使用局部 state。保存后以服务器返回对象更新列表，避免标题等字段只改在表单而未反映到页面。")
    add_heading(doc, "2.3 分页与搜索", 2)
    add_body(doc, "ApiMemoryRepository.getAll 使用游标循环请求每页数据，直到 nextCursor 为空，并记录已访问游标防止死循环。搜索覆盖标题、备忘录、心情、状态和主题。")
    add_heading(doc, "2.4 批量导入", 2)
    for item in [
        "FileList 转换为稳定的文件任务数组。",
        "统一主题时只在队列开始前执行一次长期副本策略决策。",
        "任务顺序执行，降低本地后端和磁盘瞬时压力。",
        "每项独立记录成功或失败，失败不回滚成功图片。",
        "重试只重新提交失败任务，避免重复创建成功记忆。",
    ]: add_bullet(doc, item)

    add_heading(doc, "3 后端方案", 1)
    add_heading(doc, "3.1 NestJS 模块", 2)
    add_table(doc, ["模块", "接口前缀", "核心职责"], [
        ["MediaModule", "/api/media", "上传、签名校验、受控访问和孤儿媒体清理"],
        ["MemoriesModule", "/api/memories", "创建、查询、编辑、软删除和自动关联"],
        ["TopicsModule", "/api/topics", "主题列表、自定义主题和提示抑制设置"],
        ["LongTermCopiesModule", "/api/long-term-copies", "副本 CRUD、回填和时间线"],
        ["PrismaModule", "内部模块", "PrismaClient 生命周期和事务访问"],
    ], widths=[1.65, 2.1, 2.7])
    add_heading(doc, "3.2 事务边界", 2)
    add_body(doc, "图片先通过 MediaModule 独立上传。记忆创建在单个事务中完成 Memory、MemoryTopic 和 LongTermCopyMemory 写入。编辑在同一事务中更新基础字段、重建主题关系、删除不再匹配的副本关联并补充新关联。副本创建与可选历史回填也在事务中执行；授权弹窗触发的新副本明确不执行历史回填。")
    add_heading(doc, "3.3 自动收集查询", 2)
    add_code(doc, "WHERE userId = :userId\n  AND topicId IN (:selectedTopicIds)\n  AND lifecycleStatus = 'ACTIVE'\n  AND autoCollectEnabled = true")
    add_body(doc, "只有查询返回的长期副本才建立 LongTermCopyMemory。同一用户和主题通过数据库唯一约束保证最多一个副本；并发冲突转换为 HTTP 409，前端刷新副本并复用已创建的记录。")
    add_heading(doc, "3.4 输入与媒体校验", 2)
    add_body(doc, "所有 JSON 接口使用 class-validator DTO，启用 whitelist、forbidNonWhitelisted 和 transform。布尔字段必须提交真正的 JSON boolean，字符串“false”不会被当作 true。上传在 Multer 层限制为 20MB，并在服务层同时校验 JPEG、PNG、WebP 的 MIME 与文件签名。")
    add_heading(doc, "3.5 错误语义", 2)
    add_table(doc, ["场景", "HTTP", "前端处理"], [
        ["输入缺失或主题无效", "400", "保留表单并显示原因"],
        ["记忆或副本不存在", "404", "关闭失效详情并刷新列表"],
        ["同主题副本重复", "409", "提示直接使用已有副本"],
        ["数据库或文件失败", "500", "显示重试，批量任务记录失败项"],
    ], widths=[2.15, 1.2, 3.1])

    add_heading(doc, "4 数据模型", 1)
    add_table(doc, ["实体", "关键字段", "说明"], [
        ["MediaAsset", "storagePath、mimeType、fileSize", "图片元数据和文件定位"],
        ["Memory", "title、note、mood、statuses、eventAt", "照片记忆主体"],
        ["Memory.mediaId", "mediaId 唯一外键", "当前 MVP 一条记忆对应一张图片"],
        ["Topic", "name、isNone、suppressLongTermPrompt", "主题与提示策略"],
        ["MemoryTopic", "memoryId、topicId", "记忆与主题多对多"],
        ["LongTermCopy", "topicId、status、autoCollectEnabled", "主题长期时间线"],
        ["LongTermCopyMemory", "copyId、memoryId、linkedAt", "副本记忆节点；展示时间读取 Memory.eventAt"],
    ], widths=[1.55, 3.05, 1.85])
    add_heading(doc, "4.1 数据库迁移", 2)
    add_body(doc, "0002_memory_fields_and_copy_policy 新增 Memory.title、Memory.mood、Topic.suppressLongTermPrompt 和 LongTermCopy.autoCollectEnabled，移除 Memory.isPending，并增加 LongTermCopy(userId, topicId) 唯一索引。")

    add_heading(doc, "5 API 设计", 1)
    add_table(doc, ["方法", "路径", "用途"], [
        ["GET", "/api/memories", "游标分页读取记忆，可带 keyword 和 topicId"],
        ["POST", "/api/media", "multipart 上传图片并返回 mediaId"],
        ["DELETE", "/api/media/:id", "删除未被记忆或封面引用的媒体"],
        ["POST", "/api/memories", "JSON 创建记忆并引用 mediaId"],
        ["PATCH", "/api/memories/:id", "更新标题、文字、心情、状态、主题和时间"],
        ["DELETE", "/api/memories/:id", "软删除记忆"],
        ["GET/POST", "/api/topics", "读取或创建主题"],
        ["PATCH", "/api/topics/:id", "更新主题级提示抑制设置"],
        ["GET/POST", "/api/long-term-copies", "读取或创建长期副本"],
        ["PATCH", "/api/long-term-copies/:id", "更新状态与自动收集开关"],
        ["POST", "/api/long-term-copies/:id/backfill", "回填已有同主题记忆"],
        ["GET", "/api/long-term-copies/:id/timeline", "分页读取副本时间线"],
    ], widths=[0.8, 2.75, 2.9])
    add_heading(doc, "5.1 创建记忆请求", 2)
    add_code(doc, "# 1. POST /api/media (multipart/form-data)\nfile: <binary>\n\n# 2. POST /api/memories (application/json)\n{\n  \"mediaId\": \"media-id\",\n  \"title\": \"第一次一起看日落\",\n  \"note\": \"风很大，但我们聊了很久\",\n  \"mood\": \"开心\",\n  \"statuses\": [\"约会\", \"散步\"],\n  \"topics\": [\"topic-love\"],\n  \"eventAt\": \"2026-10-12T10:36:00.000Z\"\n}")
    add_heading(doc, "5.2 响应字段", 2)
    add_body(doc, "记忆返回 id、mediaId、imageUrl、title、note、mood、statuses、topicIds、topics、eventAt、timezone、createdAt 和 updatedAt；分页响应返回 items 与 nextCursor。前端不得把 mood 混入 statuses。")

    add_heading(doc, "6 核心业务时序", 1)
    add_heading(doc, "6.1 单图保存", 2)
    for item in [
        "前端验证图片与事件时间并收集字段。",
        "根据 topics 与 copies 计算需要决策的主题。",
        "用户确认后创建或启用副本，或保存主题抑制设置。",
        "前端先上传图片取得 mediaId，再提交 JSON 记忆请求。",
        "后端事务创建记忆与主题关系，并关联符合条件的副本。",
        "前端使用服务器返回对象刷新列表。",
    ]: add_bullet(doc, item)
    add_heading(doc, "6.2 编辑与批量", 2)
    add_body(doc, "PATCH 更新后，后端删除已移除主题对应的副本关联，再为新主题执行幂等关联。若记忆保存失败，前端删除本次新上传且未被引用的媒体；替换图片成功后清理旧媒体。批量导入先处理一次统一主题授权，再逐张创建只包含图片、时间和可选主题的记忆；成功保留，失败单独重试。")

    add_heading(doc, "7 本地开发与启动", 1)
    add_table(doc, ["项目", "本地值"], [
        ["项目根目录", r"D:\yxy5.0"], ["前端端口", "5173"],
        ["后端端口", "3001"], ["PostgreSQL", "localhost:5433"],
        ["数据库", "photo_memory"], ["Prisma 配置", r"D:\yxy5.0\server\.env"],
    ], widths=[1.75, 4.7])
    add_heading(doc, "7.1 启动顺序", 2)
    add_code(doc, "# 终端一\ncd D:\\yxy5.0\\server\nnpx prisma migrate deploy\nnpx prisma db seed\nnpm run dev\n\n# 终端二\ncd D:\\yxy5.0\nnpm run dev")
    add_body(doc, "DATABASE_URL 使用 postgresql://用户名:密码@localhost:5433/photo_memory?schema=public。密码中的 @、#、% 等字符必须 URL 编码。Vite 代理目标与后端端口保持一致。")
    add_heading(doc, "7.2 图片存储", 2)
    add_body(doc, "D:\\yxy5.0\\data\\uploads 由后端创建和写入。当前实现限制 JPEG、PNG、WebP、20MB，并校验文件签名和资源所有权；生产环境应替换为私有对象存储。记忆删除当前为软删除，媒体保留以支持后续回收站；未引用媒体可通过清理接口物理删除。")

    add_heading(doc, "8 测试与质量保证", 1)
    add_table(doc, ["层级", "覆盖重点"], [
        ["领域单元测试", "搜索、长期副本候选、批量任务构造"],
        ["组件测试", "表单字段、默认开关、决策选项、空状态和批量提示"],
        ["Repository 测试", "请求映射、分页循环、标题和心情字段"],
        ["后端服务测试", "事务、自动关联、唯一冲突、稳定分页、DTO、媒体签名与清理"],
        ["构建验证", "前端 TypeScript/Vite 与后端 Nest 构建"],
        ["API 冒烟", "主题、记忆、标题编辑、副本和时间线"],
        ["人工验收", "桌面与窄屏、单图与批量、授权三种选择"],
    ], widths=[1.55, 4.9])
    add_callout_paragraph(doc, "发布门禁  ", "数据库迁移、前后端全量测试、两端生产构建和关键 API 冒烟必须全部通过；发现回归时先新增失败测试，再修复代码。")

    add_heading(doc, "9 安全与演进", 1)
    for item in [
        "本地固定用户仅用于 MVP；上线前接入认证并从鉴权上下文读取 userId。",
        "限制上传类型、大小和文件名，媒体接口校验资源所有权。",
        "日志不得输出数据库密码、原始备忘录或完整图片路径。",
        "数据库定期备份，删除策略区分软删除与物理媒体清理。",
        "总结功能新增 Summary、SummaryPermission 和 SummarySnapshot，采用授权、预览、确认保存链路。",
        "批量规模增大时再引入对象存储、Redis 和任务队列。",
    ]: add_bullet(doc, item)

    doc.save(OUTPUT); print(OUTPUT)


if __name__ == "__main__": build()
